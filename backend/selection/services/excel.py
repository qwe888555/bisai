"""Excel export for score summaries and selection results."""

from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..common import BusinessError
from ..enums import GroupType, RoundType
from . import results as result_service
from . import scores as score_service

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
NOTE_FONT = Font(italic=True, color="595959", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _timestamp():
    return timezone.localtime().strftime("%Y%m%d_%H%M%S")


def _display_width(value):
    text = "" if value is None else str(value)
    return sum(2 if ord(char) > 127 else 1 for char in text)


def _write_header(ws, headers):
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def _write_row(ws, row_index, values):
    for column, value in enumerate(values, start=1):
        cell = ws.cell(
            row=row_index,
            column=column,
            value="" if value is None else value,
        )
        cell.alignment = CENTER
        cell.border = BORDER
        if isinstance(value, float):
            cell.number_format = "0.00"


def _set_column_widths(ws, headers, rows):
    for column, header in enumerate(headers, start=1):
        width = _display_width(header)
        for row in rows:
            if column <= len(row):
                width = max(width, _display_width(row[column - 1]))
        ws.column_dimensions[get_column_letter(column)].width = min(
            max(10, width + 3), 42
        )


def _excel_response(workbook, filename):
    buffer = BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()

    response = HttpResponse(
        content,
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        "attachment; filename=\"selection.xlsx\"; "
        f"filename*=UTF-8''{quote(filename)}"
    )
    response["Content-Length"] = str(len(content))
    return response


def export_scores(event_id, group_type, score_round):
    summaries = score_service.summary(event_id, group_type, score_round)
    if not summaries:
        raise BusinessError(404, "没有可导出的评分数据")

    group_label = GroupType.label(group_type)
    round_label = RoundType.label(score_round)

    dimension_headers = []
    seen_codes = set()
    for summary in summaries:
        for item in summary["items"]:
            if item["dimension_code"] not in seen_codes:
                seen_codes.add(item["dimension_code"])
                dimension_headers.append(
                    (item["dimension_code"], item["dimension_name"])
                )

    headers = (
        ["排名", "姓名", "球队", "位置", "组别", "轮次"]
        + [name for _, name in dimension_headers]
        + ["总分", "评分时间"]
    )

    workbook = Workbook()
    ws = workbook.active
    ws.title = "评分汇总"
    _write_header(ws, headers)

    rows = []
    for summary in summaries:
        item_map = {
            item["dimension_code"]: item["score"]
            for item in summary["items"]
        }
        row = [
            summary["rank_no"],
            summary["candidate_name"],
            summary["team_name"],
            summary["position"],
            summary["group_type_label"],
            summary["round_label"],
        ]
        row += [
            item_map.get(code) for code, _ in dimension_headers
        ]
        row += [summary["total_score"], summary["updated_at"]]
        rows.append(row)
        _write_row(ws, len(rows) + 1, row)

    _set_column_widths(ws, headers, rows)
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    )

    filename = f"遴选评分_{group_label}_{round_label}_{_timestamp()}.xlsx"
    return _excel_response(workbook, filename)


def export_results(event_id, group_type):
    batch = result_service.latest_result(event_id, group_type)
    if batch["stale"]:
        raise BusinessError(
            409, "评分已变更，遴选结果已过期，请重新生成后再导出"
        )

    group_label = batch["group_type_label"]
    headers = [
        "排名",
        "姓名",
        "球队",
        "位置",
        "组别",
        "结果依据",
        "初选总分",
        "复选总分",
        "综合分",
        "生成时间",
    ]

    workbook = Workbook()
    ws = workbook.active
    ws.title = "遴选结果"
    _write_header(ws, headers)

    rows = []
    for item in batch["items"]:
        row = [
            item["rank_no"],
            item["candidate_name"],
            item["team_name"],
            item["position"],
            group_label,
            batch["result_round_label"],
            item["initial_score"],
            item["final_score"],
            item["composite_score"],
            batch["generated_at"],
        ]
        rows.append(row)
        _write_row(ws, len(rows) + 1, row)

    data_last_row = len(rows) + 1
    _set_column_widths(ws, headers, rows)
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{data_last_row}"
    )

    note_row = data_last_row + 2
    ws.cell(
        row=note_row,
        column=1,
        value=f"计算规则：{batch['formula_snapshot']}",
    )
    ws.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=len(headers),
    )
    note_cell = ws.cell(row=note_row, column=1)
    note_cell.font = NOTE_FONT
    note_cell.alignment = LEFT
    ws.row_dimensions[note_row].height = 34

    version_row = note_row + 1
    ws.cell(
        row=version_row,
        column=1,
        value=(
            f"结果批次：v{batch['version_no']}    "
            f"生成时间：{batch['generated_at']}"
        ),
    )
    ws.merge_cells(
        start_row=version_row,
        start_column=1,
        end_row=version_row,
        end_column=len(headers),
    )
    version_cell = ws.cell(row=version_row, column=1)
    version_cell.font = NOTE_FONT
    version_cell.alignment = LEFT

    filename = f"遴选结果_{group_label}_{_timestamp()}.xlsx"
    return _excel_response(workbook, filename)
