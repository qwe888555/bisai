"""批量导入、文档生成、Excel 导出与统计口径。"""

from __future__ import annotations

import io
import json
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader

from apps.enrollments import imports
from apps.enrollments.clients import person_client
from apps.enrollments.tests.factories import (
    BASE,
    add_player,
    approve_flow,
    complete_draft,
    create_draft,
    get_detail,
)

WEIGHTS = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
CHECK_CODES = ("1", "0", "X", "9", "8", "7", "6", "5", "4", "3", "2")


def make_id(prefix17: str) -> str:
    total = sum(int(prefix17[index]) * WEIGHTS[index] for index in range(17))
    return prefix17 + CHECK_CODES[total % 11]


def roster_workbook(rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(imports.TEMPLATE_HEADERS))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def post_xlsx(client, url: str, headers: dict, data: bytes, mode: str):
    return client.post(
        url,
        data={
            "mode": mode,
            "file": SimpleUploadedFile(
                "players.xlsx",
                data,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        **headers,
    )


def test_roster_template_headers(client, auth):
    headers = auth("team")
    response = client.get(f"{BASE}/imports/templates/roster.xlsx", **headers)
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    header = tuple(cell.value for cell in workbook.active[1])
    assert header == imports.TEMPLATE_HEADERS


def test_import_players_validate_then_apply(client, auth):
    headers = auth("team")
    enrollment_id = create_draft(client, headers)["id"]
    valid_id = make_id("11010520100101001")
    rows = [
        ("张三", valid_id, "SCH001", "ST001", "男", "2010-05-01", 9, "GK"),
        ("李四", "123456789012345678", "SCH001", "", "男", "2010-05-01", 10, "DF"),
    ]

    version = get_detail(client, headers, enrollment_id)["version"]
    validate = post_xlsx(
        client,
        f"{BASE}/enrollments/{enrollment_id}/imports/players?version={version}",
        headers,
        roster_workbook(rows),
        "VALIDATE_ONLY",
    )
    assert validate.status_code == 200, validate.json()
    report = validate.json()["data"]
    assert report["total"] == 2
    assert report["successCount"] == 1
    assert report["failureCount"] == 1
    assert report["errors"][0]["code"] == "INVALID_ID_NUMBER"
    assert get_detail(client, headers, enrollment_id)["players"] == []

    errors_csv = client.get(
        f"{BASE}/imports/{report['batchId']}/errors.csv", **headers
    )
    assert errors_csv.status_code == 200
    text = errors_csv.content.decode("utf-8-sig")
    assert "INVALID_ID_NUMBER" in text
    assert "123456789012345678" not in text

    version = get_detail(client, headers, enrollment_id)["version"]
    apply = post_xlsx(
        client,
        f"{BASE}/enrollments/{enrollment_id}/imports/players?version={version}",
        headers,
        roster_workbook([rows[0]]),
        "APPLY_VALID_ROWS",
    )
    assert apply.status_code == 200, apply.json()
    assert apply.json()["data"]["successCount"] == 1
    detail = get_detail(client, headers, enrollment_id)
    assert len(detail["players"]) == 1
    assert detail["players"][0]["athleteName"] == "张三"
    assert detail["version"] == version + 1


def test_photo_zip_matching_and_unmatched(client, auth):
    headers = auth("team")
    enrollment_id = create_draft(client, headers)["id"]
    matched_id = make_id("11010520100101003")
    athlete_id = person_client.resolve_person("", matched_id)["id"]
    add_player(client, headers, enrollment_id, athlete_id, jersey=7, position="GK")

    image_buffer = io.BytesIO()
    Image.new("RGB", (16, 16), (200, 10, 10)).save(image_buffer, "JPEG")
    image_bytes = image_buffer.getvalue()

    def zip_with(name: str) -> bytes:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr(name, image_bytes)
        return buffer.getvalue()

    version = get_detail(client, headers, enrollment_id)["version"]
    matched = client.post(
        f"{BASE}/enrollments/{enrollment_id}/imports/photos?version={version}",
        data={"file": SimpleUploadedFile("photos.zip", zip_with(f"{matched_id}.jpg"))},
        **headers,
    )
    assert matched.status_code == 200, matched.json()
    report = matched.json()["data"]
    assert report["mode"] == "PHOTO_IMPORT"
    assert report["successCount"] == 1
    detail = get_detail(client, headers, enrollment_id)
    assert detail["players"][0]["photoFileId"]

    unmatched_id = make_id("11010520100101004")
    version = detail["version"]
    unmatched = client.post(
        f"{BASE}/enrollments/{enrollment_id}/imports/photos?version={version}",
        data={"file": SimpleUploadedFile("photos2.zip", zip_with(f"{unmatched_id}.jpg"))},
        **headers,
    )
    assert unmatched.status_code == 200, unmatched.json()
    report = unmatched.json()["data"]
    assert report["failureCount"] == 1
    assert report["errors"][0]["code"] == "PHOTO_UNMATCHED"
    assert matched_id not in str(report["errors"])


def test_documents_bundle_and_excel_are_masked(client, auth):
    team_headers = auth("team")
    city_headers = auth("city")
    province_headers = auth("province")
    enrollment_id = complete_draft(client, team_headers)
    approved = approve_flow(
        client, team_headers, city_headers, province_headers, enrollment_id, "doc"
    )
    assert approved["status"] == "APPROVED"

    pdf = client.get(
        f"{BASE}/enrollments/{enrollment_id}/documents/registration.pdf",
        **team_headers,
    )
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"
    assert len(PdfReader(io.BytesIO(pdf.content)).pages) == 1

    docx = client.get(
        f"{BASE}/enrollments/{enrollment_id}/documents/players.docx", **team_headers
    )
    assert docx.status_code == 200
    assert docx.content[:2] == b"PK"

    export_body = {"filters": {"eventId": "event_01"}, "scope": "FILTERED"}
    bundle = client.post(
        f"{BASE}/enrollment-exports/documents.zip",
        data=json.dumps(export_body),
        content_type="application/json",
        **province_headers,
    )
    assert bundle.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(bundle.content)).namelist()
    assert len(names) == 1
    assert names[0].endswith("_报名表.pdf")

    excel = client.post(
        f"{BASE}/enrollment-exports/registrations.xlsx",
        data=json.dumps(export_body),
        content_type="application/json",
        **province_headers,
    )
    assert excel.status_code == 200
    workbook = load_workbook(io.BytesIO(excel.content))
    values = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    joined = " ".join(values)
    assert "13800000001" not in joined
    assert "138****0001" in joined


def test_statistics_excludes_draft_and_dedups(client, auth):
    team_headers = auth("team")
    other_headers = auth("team2")
    city_headers = auth("city")
    province_headers = auth("province")

    approved_id = complete_draft(client, team_headers)
    approve_flow(
        client, team_headers, city_headers, province_headers, approved_id, "stat"
    )
    # 另一队的草稿不应计入
    complete_draft(
        client, other_headers, team_id="org_team_02", players=2, materials=False
    )

    response = client.get(
        f"{BASE}/enrollment-statistics?eventId=event_01&dimension=CITY",
        **province_headers,
    )
    assert response.status_code == 200
    rows = response.json()["data"]["items"]
    assert len(rows) == 1
    assert rows[0]["key"] == "org_city_01"
    assert rows[0]["teamCount"] == 1
    assert rows[0]["athleteCount"] == 5
    assert rows[0]["approvedCount"] == 1

    by_group = client.get(
        f"{BASE}/enrollment-statistics?eventId=event_01&dimension=GROUP",
        **province_headers,
    )
    assert by_group.status_code == 200
    assert by_group.json()["data"]["items"][0]["teamCount"] == 1
