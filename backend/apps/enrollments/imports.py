"""批量导入：球员 Excel 模板/校验/写入、照片 ZIP 匹配、导入错误 CSV。

同步处理，限制行数与体积；逐行原子；证件号只用于匹配，不进入日志/响应。
"""

from __future__ import annotations

import csv
import io
import os
import re
import zipfile
from datetime import date
from typing import Optional

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.utils import timezone

from apps.common.crypto import mask_filename, mask_id_number, normalize_id_number
from apps.common.errors import (
    BATCH_LIMIT_EXCEEDED,
    FILE_TOO_LARGE,
    INVALID_ARGUMENT,
    STATE_CONFLICT,
    UNSUPPORTED_FILE_TYPE,
    ApiError,
)
from apps.common.ids import new_id
from apps.common.versioning import bump, check_version

from . import scope
from .clients import file_client, person_client, team_client
from .models import (
    EDITABLE_STATUSES,
    QUALIFICATION_EXCLUDED,
    RosterItem,
    Enrollment,
    ImportBatch,
    ImportError,
)

TEMPLATE_HEADERS = (
    "姓名",
    "身份证号",
    "学校编码",
    "学籍号",
    "性别",
    "出生日期",
    "球衣号码",
    "位置",
)
MAX_ROWS = 1000
MAX_ZIP_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_UNCOMPRESSED = 100 * 1024 * 1024

_WEIGHTS = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
_CHECK_CODES = ("1", "0", "X", "9", "8", "7", "6", "5", "4", "3", "2")

_GENDER_MAP = {
    "男": "MALE",
    "MALE": "MALE",
    "女": "FEMALE",
    "FEMALE": "FEMALE",
    "未指定": "UNSPECIFIED",
    "UNSPECIFIED": "UNSPECIFIED",
}


def valid_id_number(value: str) -> bool:
    """18 位身份证 + ISO 7064 MOD 11-2 校验位。"""
    text = normalize_id_number(value)
    if not re.fullmatch(r"\d{17}[\dX]", text):
        return False
    total = sum(int(text[index]) * _WEIGHTS[index] for index in range(17))
    return _CHECK_CODES[total % 11] == text[17]


# ---------------------------------------------------------------- 模板


def render_roster_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "球员名单"
    sheet.append(list(TEMPLATE_HEADERS))
    fill = PatternFill("solid", fgColor="D9D9D9")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    # 身份证号/学籍号按文本列写入，避免科学计数法
    for row in range(2, MAX_ROWS + 2):
        sheet.cell(row=row, column=2).number_format = "@"
        sheet.cell(row=row, column=4).number_format = "@"
    widths = (14, 22, 14, 20, 10, 14, 12, 14)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- 球员导入


def _read_rows(uploaded) -> tuple[list[tuple[int, list[str]]], list[dict]]:
    name = (uploaded.name or "").lower()
    if not name.endswith(".xlsx"):
        raise ApiError(
            UNSUPPORTED_FILE_TYPE,
            "球员名单仅支持 xlsx",
            [{"field": "file", "reason": "需要 .xlsx"}],
        )
    from openpyxl import load_workbook

    data = uploaded.read()
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ApiError(
            UNSUPPORTED_FILE_TYPE,
            "无法解析该 xlsx 文件",
            [{"field": "file", "reason": "文件损坏或格式不符"}],
        ) from exc
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ApiError(INVALID_ARGUMENT, "模板为空")
    header = tuple(str(cell or "").strip() for cell in rows[0][: len(TEMPLATE_HEADERS)])
    if header != TEMPLATE_HEADERS:
        raise ApiError(
            INVALID_ARGUMENT,
            "表头与固定模板不一致，不支持 AI 表头匹配",
            [{"field": "file", "reason": "表头必须为：" + " | ".join(TEMPLATE_HEADERS)}],
        )
    data_rows: list[tuple[int, list[str]]] = []
    for index, raw in enumerate(rows[1:], start=2):
        values = ["" if cell is None else str(cell).strip() for cell in raw[: len(TEMPLATE_HEADERS)]]
        if not any(values):
            continue
        data_rows.append((index, values))
    if len(data_rows) > MAX_ROWS:
        raise ApiError(
            BATCH_LIMIT_EXCEEDED,
            f"单次最多导入 {MAX_ROWS} 行，请拆分文件",
            [{"field": "file", "reason": "行数超限"}],
        )
    return data_rows, []


def _validate_row(
    actor,
    enrollment: Enrollment,
    team: dict,
    values: list[str],
    seen_ids: set[str],
    seen_jerseys: set[int],
) -> tuple[Optional[dict], list[dict]]:
    errors: list[dict] = []
    name, id_number, school_code, student_number, gender, birth_date, jersey, position = values

    if not name:
        errors.append(_row_error("NAME_REQUIRED", "姓名必填", "姓名"))
    if not valid_id_number(id_number):
        errors.append(
            _row_error(
                "INVALID_ID_NUMBER",
                f"身份证号格式或校验位不合法：{mask_id_number(id_number)}",
                "身份证号",
            )
        )
    normalized_id = normalize_id_number(id_number)
    if normalized_id and normalized_id in seen_ids:
        errors.append(
            _row_error(
                "DUPLICATE_ID_NUMBER",
                f"文件内身份证重复：{mask_id_number(normalized_id)}",
                "身份证号",
            )
        )

    if not school_code:
        errors.append(_row_error("SCHOOL_CODE_REQUIRED", "学校编码必填", "学校编码"))
    else:
        org = _find_school(actor, team, school_code)
        if org is None:
            errors.append(
                _row_error(
                    "SCHOOL_CODE_INVALID",
                    "学校编码不存在或不在当前报送主体范围内",
                    "学校编码",
                )
            )

    gender_value = _GENDER_MAP.get((gender or "").strip().upper()) or _GENDER_MAP.get(
        (gender or "").strip()
    )
    if gender and not gender_value:
        errors.append(_row_error("INVALID_GENDER", "性别不在模板字典内", "性别"))

    parsed_date = _parse_birth_date(birth_date)
    if birth_date and parsed_date is None:
        errors.append(_row_error("INVALID_BIRTH_DATE", "出生日期格式应为 YYYY-MM-DD", "出生日期"))

    jersey_value = None
    if jersey:
        try:
            jersey_value = int(jersey)
            if jersey_value < 0 or jersey_value > 999:
                raise ValueError
        except ValueError:
            errors.append(_row_error("INVALID_JERSEY_NUMBER", "球衣号码应为 0—999 整数", "球衣号码"))
            jersey_value = None
    if jersey_value is not None:
        if jersey_value in seen_jerseys:
            errors.append(
                _row_error("DUPLICATE_JERSEY_NUMBER", f"球衣号码 {jersey_value} 在文件内重复", "球衣号码")
            )

    if position and len(position) > 50:
        errors.append(_row_error("INVALID_POSITION", "位置不超过 50 字符", "位置"))

    if errors:
        return None, errors

    seen_ids.add(normalized_id)
    if jersey_value is not None:
        seen_jerseys.add(jersey_value)
    person = person_client.resolve_person(
        name,
        normalized_id,
        student_number=student_number,
        gender=gender_value or "",
        birth_date=parsed_date.isoformat() if parsed_date else "",
        school_code=school_code,
    )
    existing = RosterItem.objects.filter(
        enrollment_id=enrollment.pk, athlete_id=person["id"]
    ).exists()
    if existing:
        return None, [
            _row_error("DUPLICATE_PLAYER", "该运动员已在本报名名单内", "身份证号")
        ]
    return (
        {
            "athlete_id": person["id"],
            "athlete_name": person.get("name") or name,
            "masked_id_number": person.get("maskedIdNumber") or mask_id_number(normalized_id),
            "jersey_number": jersey_value,
            "position": position,
            "student_status_file_ids": [],
        },
        [],
    )


def _parse_birth_date(value: str) -> Optional[date]:
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _find_school(actor, team: dict, school_code: str):
    from apps.accounts.models import Organization
    from apps.common.scopes import org_scope_ids

    org = Organization.objects.filter(code=school_code, enabled=True).first()
    if org is None:
        return None
    if scope.is_org_scoped(actor):
        return org if org.pk in set(org_scope_ids(actor)) else None
    # 球队账号只能导入本校编码
    team_org_id = team.get("orgId") or ""
    return org if org.pk == team_org_id else None


def _row_error(code: str, message: str, field: str) -> dict:
    return {"code": code, "message": message[:500], "field": field[:100]}


def _save_errors(batch_id: str, row_number: Optional[int], errors: list[dict]) -> None:
    for error in errors:
        ImportError.objects.create(
            id=new_id("ime"),
            batch_id=batch_id,
            row_number=row_number,
            field=error.get("field", ""),
            code=error["code"],
            message=error["message"],
            file_name=mask_filename(error.get("fileName", ""))[:255],
        )


@transaction.atomic
def import_players(actor, enrollment_id: str, version: int, mode: str, uploaded) -> dict:
    enrollment = (
        Enrollment.objects.select_for_update().filter(pk=enrollment_id).first()
    )
    if enrollment is None:
        from apps.common.errors import NOT_FOUND

        raise ApiError(NOT_FOUND)
    scope.ensure_enrollment_access(actor, enrollment)
    if enrollment.status not in EDITABLE_STATUSES:
        raise ApiError(STATE_CONFLICT, "仅可编辑状态允许导入")
    check_version(enrollment, version)
    from . import precheck

    precheck.ensure_window_open(enrollment.event_id)

    team = team_client.get_team(enrollment.team_id)
    rows, _ = _read_rows(uploaded)

    batch = ImportBatch.objects.create(
        id=new_id("imb"),
        enrollment_id=enrollment.pk,
        mode=mode,
        total=len(rows),
        created_by=actor.pk,
    )

    seen_ids: set[str] = set()
    seen_jerseys: set[int] = set()
    existing_jerseys = set(
        RosterItem.objects.filter(
            enrollment_id=enrollment.pk,
            enrollment__status__in=["DRAFT", "REJECTED", "WITHDRAWN"],
        )
        .exclude(qualification_status=QUALIFICATION_EXCLUDED)
        .exclude(jersey_number=None)
        .values_list("jersey_number", flat=True)
    )
    seen_jerseys.update(int(value) for value in existing_jerseys)

    success = 0
    failures: list[tuple[int, list[dict]]] = []
    prepared: list[dict] = []
    for row_number, values in rows:
        record, errors = _validate_row(
            actor, enrollment, team, values, seen_ids, seen_jerseys
        )
        if errors:
            failures.append((row_number, errors))
            continue
        prepared.append(record)
        success += 1

    if mode == "APPLY_VALID_ROWS" and prepared:
        for record in prepared:
            RosterItem.objects.create(
                id=new_id("rst"),
                enrollment_id=enrollment.pk,
                **record,
            )
        bump(enrollment)
        enrollment.save()

    batch.success_count = success
    batch.failure_count = len(failures)
    batch.save()
    for row_number, errors in failures:
        _save_errors(batch.pk, row_number, errors)

    return {
        "batchId": batch.pk,
        "mode": mode,
        "total": len(rows),
        "successCount": success,
        "failureCount": len(failures),
        "enrollmentVersion": int(enrollment.version or 0),
        "errors": [
            {
                "rowNumber": row_number,
                "field": error.get("field", ""),
                "code": error["code"],
                "message": error["message"],
                "fileName": None,
            }
            for row_number, errors in failures
            for error in errors
        ],
    }


# ---------------------------------------------------------------- 照片导入


def _safe_image_name(name: str) -> bool:
    if not name or name.endswith("/"):
        return False
    if ".." in name.replace("\\", "/").split("/"):
        return False
    if name.startswith("/") or name.startswith("\\") or re.match(r"^[A-Za-z]:", name):
        return False
    extension = os.path.splitext(name)[1].lower()
    return extension in {".jpg", ".jpeg", ".png"}


@transaction.atomic
def import_photos(actor, enrollment_id: str, version: int, uploaded) -> dict:
    enrollment = (
        Enrollment.objects.select_for_update().filter(pk=enrollment_id).first()
    )
    if enrollment is None:
        from apps.common.errors import NOT_FOUND

        raise ApiError(NOT_FOUND)
    scope.ensure_enrollment_access(actor, enrollment)
    if enrollment.status not in EDITABLE_STATUSES:
        raise ApiError(STATE_CONFLICT, "仅可编辑状态允许导入照片")
    check_version(enrollment, version)
    from . import precheck

    precheck.ensure_window_open(enrollment.event_id)

    if uploaded.size and uploaded.size > MAX_ZIP_BYTES:
        raise ApiError(FILE_TOO_LARGE, "照片 ZIP 不得超过 20 MiB")
    payload = uploaded.read()
    if len(payload) > MAX_ZIP_BYTES:
        raise ApiError(FILE_TOO_LARGE, "照片 ZIP 不得超过 20 MiB")
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as exc:
        raise ApiError(
            UNSUPPORTED_FILE_TYPE,
            "不是有效的 ZIP 文件",
            [{"field": "file", "reason": "ZIP 解析失败"}],
        ) from exc

    entries = [info for info in archive.infolist() if not info.is_dir()]
    if len(entries) > MAX_ZIP_ENTRIES:
        raise ApiError(BATCH_LIMIT_EXCEEDED, f"ZIP 内最多 {MAX_ZIP_ENTRIES} 张照片")
    total_size = sum(info.file_size for info in entries)
    if total_size > MAX_UNCOMPRESSED:
        raise ApiError(BATCH_LIMIT_EXCEEDED, "ZIP 解压后总量不得超过 100 MiB")

    batch = ImportBatch.objects.create(
        id=new_id("imb"),
        enrollment_id=enrollment.pk,
        mode="PHOTO_IMPORT",
        total=len(entries),
        created_by=actor.pk,
    )

    roster_by_person: dict[str, RosterItem] = {
        item.athlete_id: item for item in enrollment.players.all()
    }
    success = 0
    failures: list[tuple[Optional[int], list[dict]]] = []
    matched_any = False

    for info in entries:
        display_name = mask_filename(os.path.basename(info.filename))
        if not _safe_image_name(info.filename):
            failures.append(
                (
                    None,
                    [
                        {
                            "code": "UNSAFE_OR_UNSUPPORTED_ENTRY",
                            "message": "拒绝路径穿越/嵌套压缩/非图片文件",
                            "field": "file",
                            "fileName": display_name,
                        }
                    ],
                )
            )
            continue
        stem = os.path.splitext(os.path.basename(info.filename))[0]
        normalized = normalize_id_number(stem)
        if not valid_id_number(normalized):
            failures.append(
                (
                    None,
                    [
                        {
                            "code": "INVALID_PHOTO_NAME",
                            "message": "文件名不是合法身份证号（已脱敏）",
                            "field": "file",
                            "fileName": display_name,
                        }
                    ],
                )
            )
            continue
        data = archive.read(info)
        try:
            from PIL import Image

            with Image.open(io.BytesIO(data)) as image:
                image_format = (image.format or "").upper()
            if image_format not in {"JPEG", "PNG"}:
                raise ValueError("非 JPEG/PNG")
        except Exception:  # noqa: BLE001
            failures.append(
                (
                    None,
                    [
                        {
                            "code": "UNSUPPORTED_FILE_TYPE",
                            "message": "照片内容不是 JPEG/PNG",
                            "field": "file",
                            "fileName": display_name,
                        }
                    ],
                )
            )
            continue

        person = person_client.resolve_person("", normalized)
        item = roster_by_person.get(person["id"])
        if item is None:
            failures.append(
                (
                    None,
                    [
                        {
                            "code": "PHOTO_UNMATCHED",
                            "message": "未匹配到本报名名单内的运动员",
                            "field": "file",
                            "fileName": display_name,
                        }
                    ],
                )
            )
            continue
        extension = os.path.splitext(info.filename)[1].lower()
        uploaded_file = SimpleUploadedFile(
            f"photo{extension}",
            data,
            content_type="image/jpeg" if extension in {".jpg", ".jpeg"} else "image/png",
        )
        record = file_client.upload(
            actor, "ENROLLMENT", enrollment.pk, "ROSTER_PHOTO", uploaded_file
        )
        item.photo_file_id = record.pk
        bump(item)
        item.save()
        success += 1
        matched_any = True

    if matched_any:
        bump(enrollment)
        enrollment.save()

    batch.success_count = success
    batch.failure_count = len(failures)
    batch.save()
    for row_number, errors in failures:
        _save_errors(batch.pk, row_number, errors)

    return {
        "batchId": batch.pk,
        "mode": "PHOTO_IMPORT",
        "total": len(entries),
        "successCount": success,
        "failureCount": len(failures),
        "enrollmentVersion": int(enrollment.version or 0),
        "errors": [
            {
                "rowNumber": row_number,
                "field": error.get("field", ""),
                "code": error["code"],
                "message": error["message"],
                "fileName": error.get("fileName") or None,
            }
            for row_number, errors in failures
            for error in errors
        ],
    }


# ---------------------------------------------------------------- 错误清单


def import_errors_csv(actor, batch_id: str) -> tuple[bytes, str]:
    batch = ImportBatch.objects.filter(pk=batch_id).first()
    if batch is None:
        from apps.common.errors import NOT_FOUND

        raise ApiError(NOT_FOUND)
    enrollment = Enrollment.objects.filter(pk=batch.enrollment_id).first()
    if enrollment is None:
        from apps.common.errors import NOT_FOUND

        raise ApiError(NOT_FOUND)
    scope.ensure_enrollment_access(actor, enrollment)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["行号", "字段", "错误码", "说明", "文件名（脱敏）"])
    for error in batch.errors.all().order_by("row_number", "id"):
        writer.writerow(
            [
                error.row_number or "",
                error.field,
                error.code,
                error.message,
                mask_filename(error.file_name),
            ]
        )
    return buffer.getvalue().encode("utf-8-sig"), "导入错误清单.csv"
